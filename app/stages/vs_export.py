"""
Stage 6 — VoteSource export builder.

What it does, in plain English:
  Build the file we send to the CCHQ Insights team for voter-intention
  overlay. Column shape matches the canonical `vs example.xlsx` exactly:
  13 identification columns (A–M), a blank spacer at column N, and the
  Insights team fills in 17 return columns (O–AE) before sending it back.

  Only rows the classifier marked Y or T are included — no point asking
  Insights to vote-score people we've already ruled out as wrong matches.

What's different from the old process:
  Replaces Step 10 of the 18-step flow. Used to be: hand-build a
  spreadsheet ad-hoc, send as an attachment, hope Insights' intake
  script recognised the column shape. This module produces a file
  in the agreed canonical layout, every time.
"""
import os
import openpyxl
from openpyxl.styles import Font
import pandas as pd


# (master_column_name, vs_column_header). Order = left-to-right in the
# output sheet. Underscore-prefixed master keys are derived (built below).
SEND_COLUMNS = [
    ("Unique ID",                  "UniqueID"),
    ("First Name",                 "FirstName"),
    ("Middle Names",               "MiddleNames"),
    ("Surname",                    "Lastname"),
    ("Officer date of birth",      "DateOfBirth"),
    ("Officer address line one",   "OfficerAddressLineOne"),
    ("Officer address locality",   "OfficerAddressLocality"),
    ("Officer address country",    "OfficerAddressCountry"),
    ("Officer address post code",  "OfficerAddressPostcode"),
    ("Company Name",               "CompanyName"),
    ("_FullAddressWithCompany",    "FullAddressWithCompany"),
    ("_FullAddressWithoutCompany", "FullAddressWithoutCompany"),
    ("Apollo Email",               "Email"),
]


def _build_full_address(row, include_company: bool) -> str:
    parts = []
    if include_company:
        cn = str(row.get("Company Name", "")).strip()
        if cn:
            parts.append(cn)
    for col in ("Officer address line one",
                "Officer address locality",
                "Officer address post code"):
        val = str(row.get(col, "")).strip()
        if val:
            parts.append(val)
    return " ".join(parts)


def build_vs_export(project_dir, region_code, progress_cb=None):
    """
    Build the VoteSource export XLSX for Y/T rows.

    Returns:
        str — absolute path to the generated XLSX
    """
    def log(msg):
        if progress_cb:
            progress_cb(msg)

    classified_path = os.path.join(project_dir,
                                   f"master_{region_code}_classified.csv")
    if not os.path.exists(classified_path):
        raise FileNotFoundError(
            f"master_{region_code}_classified.csv not found — run Stage 5 first."
        )

    master = pd.read_csv(classified_path, dtype=str, keep_default_na=False)
    log(f"Loaded classified master: {len(master):,} rows")

    # Filter to Y/T only
    if "Result" in master.columns:
        before = len(master)
        master = master[master["Result"].isin(["Y", "T"])].copy()
        log(f"  Filtered to Y/T: {len(master):,} rows "
            f"(dropped {before - len(master):,} N rows)")

    # Build the two derived address columns
    master["_FullAddressWithCompany"] = master.apply(
        lambda r: _build_full_address(r, include_company=True), axis=1
    )
    master["_FullAddressWithoutCompany"] = master.apply(
        lambda r: _build_full_address(r, include_company=False), axis=1
    )
    log("  Built FullAddressWithCompany / FullAddressWithoutCompany")

    # ── Write XLSX ────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Row 1 headers (columns A–M). Column N (index 14) is left without a
    # header on purpose — it's the spacer between send and return cols.
    for col_idx, (_, header) in enumerate(SEND_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    # Data rows
    for r_idx, (_, row) in enumerate(master.iterrows(), start=2):
        for c_idx, (src_col, _) in enumerate(SEND_COLUMNS, start=1):
            val = row.get(src_col, "")
            ws.cell(row=r_idx, column=c_idx,
                    value=val if val != "" else None)

    out_path = os.path.join(project_dir, f"vs_export_{region_code}.xlsx")
    wb.save(out_path)
    log(f"  Saved: {os.path.basename(out_path)} "
        f"({len(master):,} rows × {len(SEND_COLUMNS)} cols)")
    return out_path
